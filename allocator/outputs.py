"""Everything the office needs after an allocation, in the shapes it uses today.

    allocation.xlsx      the master overview (all students, per project, summary, warnings)
    mail_merge.csv       one row per running project: supervisors + the group's emails
    sap/PROxxxx.xlsx     one workbook per SAP code, ready to paste into the mass booking
                         and the individual-work title import (title wrapped at 45 chars)
    canvas_groups.csv    the Canvas group import for the PRO1000 course
    lab_overview.xlsx    lab-based projects with their lab needs and actual numbers (DUB30)
    norm_hours.xlsx      every running project with its actual number (onderwijsverrekening)
    bible_updated.xlsx   the bible with "actual nr" filled in
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .model import LAB_COLUMNS, Result
from .solver import format_summary, summarise

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="001C3D")
SUB = Font(bold=True)


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        w = (widths or {}).get(h, max(10, min(60, len(str(h)) + 4)))
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def sap_code(student, project) -> str:
    """Which SAP module code this allocation books on.

    1000-level: PRO1002. 2000-level: PRO2001 for the first 2000-level project,
    PRO2002 for the second, and so on; same for 3000. ESD confirms the count
    against SAP's own red flags, this is the starting point.
    """
    if project.level == "1000":
        return "PRO1002"
    if project.level == "2000":
        return f"PRO200{min(student.completed_2000 + 1, 9)}"
    return f"PRO300{min(student.completed_3000 + 1, 9)}"


def wrap_title(title: str, width: int = 45) -> str:
    """SAP's individual-work field takes lines of at most 45 characters."""
    words, lines, cur = title.split(), [], ""
    for w in words:
        if len(cur) + len(w) + (1 if cur else 0) <= width:
            cur = f"{cur} {w}".strip()
        else:
            if cur:
                lines.append(cur)
            while len(w) > width:
                lines.append(w[:width])
                w = w[width:]
            cur = w
    if cur:
        lines.append(cur)
    return "\n".join(lines)


def write_all(result: Result, projects, students, preferences, out_dir, *, period="", year="",
              session="") -> list[str]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    proj = {p.code: p for p in projects}
    stud = {s.id: s for s in students}
    pref = {p.student_id: p for p in preferences}
    by_student = result.by_student()
    by_project = result.by_project()
    written = []

    # ---- allocation.xlsx ---------------------------------------------------
    wb = Workbook()
    wb.remove(wb.active)
    rows = []
    for s in sorted(students, key=lambda s: (by_student[s.id].code or "zzz", s.name, s.id)):
        a = by_student[s.id]
        p = proj.get(a.code) if a.code else None
        rows.append([
            s.id, s.name, s.batch, a.code or "", p.title if p else "", p.level if p else "",
            p.supervisor if p else "", "; ".join(p.supervisor_emails) if p else "",
            a.rank if a.rank else "", a.reason,
            ", ".join(pref[s.id].choices) if s.id in pref else "(no form)",
            sap_code(s, p) if p else "",
        ])
    _sheet(wb, "All students",
           ["Student ID", "Name", "Batch", "Project", "Title", "Level", "Supervisor", "Supervisor emails",
            "Choice rank", "How placed", "Choices submitted", "SAP code"], rows,
           {"Title": 50, "Name": 28, "Supervisor": 24, "Supervisor emails": 40, "Choices submitted": 24})

    ws = wb.create_sheet("By project")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 60
    r = 1
    for code in sorted(proj):
        p = proj[code]
        members = by_project.get(code, [])
        ws.cell(r, 1, code).font = SUB
        ws.cell(r, 2, p.title).font = SUB
        ws.cell(r, 3, f"{p.level}-level  |  {p.supervisor}{' + ' + p.co_supervisor if p.co_supervisor else ''}"
                      f"  |  min {p.min_students} max {p.max_students}  |  placed {len(members)}"
                      f"{'  |  NOT RUNNING' if not members and p.status == 'open' else ''}"
                      f"{'  |  CLOSED' if p.status == 'closed' else ''}").font = SUB
        r += 1
        for sid in sorted(members, key=lambda i: stud[i].name):
            a = by_student[sid]
            ws.cell(r, 1, sid)
            ws.cell(r, 2, stud[sid].name)
            ws.cell(r, 3, a.reason)
            r += 1
        r += 1

    summary = summarise(result, projects, students, preferences)
    text = format_summary(summary, projects)
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 120
    for i, line in enumerate(text.splitlines(), 1):
        ws.cell(i, 1, line)
    if result.warnings:
        _sheet(wb, "Warnings", ["Warning"], [[w] for w in result.warnings], {"Warning": 120})
    unassigned = [[sid, stud[sid].name, stud[sid].batch, ", ".join(pref[sid].choices) if sid in pref else "(no form)"]
                  for sid in result.unassigned]
    if unassigned:
        _sheet(wb, "Unassigned", ["Student ID", "Name", "Batch", "Choices"], unassigned, {"Name": 28, "Choices": 30})
    f = out / "allocation.xlsx"
    wb.save(f)
    written.append(str(f))
    (out / "summary.txt").write_text(text + "\n", encoding="utf-8")
    written.append(str(out / "summary.txt"))

    # ---- mail_merge.csv ----------------------------------------------------
    f = out / "mail_merge.csv"
    with open(f, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["code", "level", "title", "supervisor", "co_supervisor", "supervisor_emails",
                    "n_students", "student_names", "student_emails", "student_ids"])
        for code in sorted(by_project):
            p = proj[code]
            members = sorted(by_project[code], key=lambda i: stud[i].name)
            w.writerow([code, p.level, p.title, p.supervisor, p.co_supervisor, "; ".join(p.supervisor_emails),
                        len(members), "; ".join(stud[i].name for i in members),
                        "; ".join(stud[i].email for i in members if stud[i].email),
                        "; ".join(members)])
    written.append(str(f))

    # ---- SAP per code --------------------------------------------------------
    sap_dir = out / "sap"
    sap_dir.mkdir(exist_ok=True)
    per_code: dict[str, list] = {}
    for s in students:
        a = by_student[s.id]
        if a.code:
            per_code.setdefault(sap_code(s, proj[a.code]), []).append((s, proj[a.code]))
    for code, items in sorted(per_code.items()):
        wb = Workbook()
        ws = wb.active
        ws.title = code
        ws.append(["Student", "Student Name", "Module", "Year", "Period", "SubType", "Individual Work", "Project"])
        for c in ws[1]:
            c.font, c.fill = HEAD, HEAD_FILL
        for s, p in sorted(items, key=lambda t: int(t[0].id)):
            ws.append([int(s.id), s.name, code, year, session, "9714", wrap_title(p.title), p.code])
            ws.cell(ws.max_row, 7).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["G"].width = 48
        f = sap_dir / f"{code}.xlsx"
        wb.save(f)
        written.append(str(f))
        with open(sap_dir / f"{code}_ids.txt", "w", encoding="utf-8") as fh:   # paste into ZPIQ_MASS_BOOKING
            fh.write("\n".join(s.id for s, _ in sorted(items, key=lambda t: int(t[0].id))) + "\n")

    # ---- Canvas groups -------------------------------------------------------
    f = out / "canvas_groups.csv"
    with open(f, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["name", "login_id", "sections", "group_name"])
        for s in sorted(students, key=lambda s: s.name):
            a = by_student[s.id]
            if a.code:
                w.writerow([s.name, s.login_id, "Section PRO1000", proj[a.code].label[:80]])
    written.append(str(f))

    # ---- lab overview + norm hours + bible ------------------------------------
    counts = result.counts
    lab_rows, norm_rows, bible_rows = [], [], []
    for p in sorted(projects, key=lambda p: p.code):
        n = counts.get(p.code, 0)
        base = [p.level, p.code, p.title, p.supervisor, p.supervisor_email, p.slp_leaders, p.co_supervisor,
                p.co_supervisor_email, p.min_students, p.max_students, n]
        bible_rows.append(base + [p.location, *[p.labs.get(l, "") for l in LAB_COLUMNS], p.equipment,
                                  p.fume_hoods, p.remarks, "not running" if n == 0 else ""])
        if n:
            norm_rows.append(base)
            if p.lab_based:
                lab_rows.append([p.code, p.level, p.title, p.supervisor, p.co_supervisor, n, p.location,
                                 *[p.labs.get(l, "") for l in LAB_COLUMNS], p.equipment, p.fume_hoods, p.remarks])
    head_base = ["Level", "Course Code", "Project title", "Supervisor name", "Supervisor e-mail", "SLP",
                 "2nd (co-)supervisor", "Co-supervisor e-mail", "min", "max", "actual nr"]
    wb = Workbook(); wb.remove(wb.active)
    _sheet(wb, "DUB30", ["Code", "Level", "Project title", "Supervisor", "Co-supervisor", "Students", "lab/ non-lab",
                         *LAB_COLUMNS, "equipment", "Fume hoods", "Remarks"], lab_rows,
           {"Project title": 50, "Supervisor": 24, "equipment": 40, "Remarks": 30})
    f = out / "lab_overview.xlsx"; wb.save(f); written.append(str(f))
    wb = Workbook(); wb.remove(wb.active)
    _sheet(wb, "Norm hours", head_base, norm_rows, {"Project title": 50, "Supervisor name": 24})
    f = out / "norm_hours.xlsx"; wb.save(f); written.append(str(f))
    wb = Workbook(); wb.remove(wb.active)
    _sheet(wb, "Sheet1", head_base + ["lab/ non-lab-based", *LAB_COLUMNS, "equipment", "Fume hoods", "Remarks", "Status"],
           bible_rows, {"Project title": 50, "Supervisor name": 24, "equipment": 36})
    f = out / "bible_updated.xlsx"; wb.save(f); written.append(str(f))
    return written
